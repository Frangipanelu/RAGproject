"""
FastAPI 后端 - 暴露 REST API
复用现有业务逻辑：workspace_manager, knowledge_base, rag, answer_saver
"""
import os
import time
import json
import shutil
import hashlib
from io import BytesIO
from pathlib import Path
from typing import Optional, List

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import config_data as config
from file_history_store import get_history
from knowledge_base import (
    KnowledgeBaseService, list_documents, get_all_tags,
    delete_document, auto_categorize, extract_tags, load_meta_db,
    save_meta_db, get_document,
)
from rag import RagService
from workspace_manager import (
    create_workspace, list_workspaces, get_workspace, delete_workspace,
    create_category, list_categories, get_category, get_category_tree,
    delete_category, update_category, get_all_descendants,
)
from answer_saver import save_answer_to_kb, build_answer_document

UPLOAD_ROOT = Path("data") / "uploads"
STATIC_DIR = Path("static")
ALLOWED_EXTS = {"md", "txt", "pdf", "docx", "html", "csv", "json", "yaml", "yml"}

# 内存缓存 RAG 实例
_RAG_CACHE: dict = {}
_KB_SERVICE = None


def get_kb_service() -> KnowledgeBaseService:
    global _KB_SERVICE
    if _KB_SERVICE is None:
        _KB_SERVICE = KnowledgeBaseService()
    return _KB_SERVICE


def get_rag_service(workspace_id: str, category_id: Optional[str] = None) -> RagService:
    key = f"{workspace_id}_{category_id or 'all'}"
    if key not in _RAG_CACHE:
        _RAG_CACHE[key] = RagService(workspace_id=workspace_id, category_id=category_id)
    return _RAG_CACHE[key]


def extract_text(file_bytes: bytes, file_ext: str) -> str:
    if file_ext in ("md", "txt", "yaml", "yml"):
        return file_bytes.decode("utf-8", errors="ignore")
    if file_ext == "pdf":
        from pypdf import PdfReader
        pdf = PdfReader(BytesIO(file_bytes))
        return "\n".join((p.extract_text() or "") for p in pdf.pages)
    if file_ext == "docx":
        from docx import Document
        doc = Document(BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)
    if file_ext == "html":
        from html.parser import HTMLParser
        class TE(HTMLParser):
            def __init__(self): super().__init__(); self.t=[]; self.s=False
            def handle_starttag(self,t,a):
                if t in ("script","style"): self.s=True
            def handle_endtag(self,t):
                if t in ("script","style"): self.s=False
            def handle_data(self,d):
                if not self.s: self.t.append(d)
        e=TE(); e.feed(file_bytes.decode("utf-8", errors="ignore"))
        return "\n".join(e.t)
    if file_ext == "csv":
        import csv
        lines = file_bytes.decode("utf-8", errors="ignore").splitlines()
        return "\n".join(" | ".join(r) for r in csv.reader(lines))
    if file_ext == "json":
        return json.dumps(json.loads(file_bytes.decode("utf-8")), ensure_ascii=False, indent=2)
    raise ValueError(f"不支持: {file_ext}")


def save_upload_bytes(file_bytes: bytes, file_name: str, file_ext: str) -> Path:
    type_dir = UPLOAD_ROOT / file_ext
    type_dir.mkdir(parents=True, exist_ok=True)
    save_path = type_dir / file_name
    base, suf = Path(file_name).stem, Path(file_name).suffix
    i = 1
    while save_path.exists():
        save_path = type_dir / f"{base}_{i}{suf}"
        i += 1
    save_path.write_bytes(file_bytes)
    return save_path


app = FastAPI(title="Knowledge Base API", version="2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


# ───────────── Pydantic Models ─────────────

class WorkspaceCreate(BaseModel):
    name: str
    description: str = ""
    icon: str = "📚"


class CategoryCreate(BaseModel):
    name: str
    workspace_id: str
    parent_id: Optional[str] = None
    icon: str = "📁"


class CategoryUpdate(BaseModel):
    name: Optional[str] = None
    icon: Optional[str] = None


class NoteCreate(BaseModel):
    title: str
    content: str
    workspace_id: str
    category_id: Optional[str] = None
    tags: List[str] = []


class ChatRequest(BaseModel):
    question: str
    workspace_id: str
    category_id: Optional[str] = None
    session_id: str = "default"


class SaveAnswerRequest(BaseModel):
    question: str
    answer: str
    workspace_id: str
    category_id: Optional[str] = None
    title: str
    sources: List[dict] = []


# ───────────── Workspace API ─────────────

@app.get("/api/workspaces")
def api_list_workspaces():
    return {"workspaces": list_workspaces()}


@app.post("/api/workspaces")
def api_create_workspace(req: WorkspaceCreate):
    ws_id = create_workspace(req.name, req.description, req.icon)
    return {"id": ws_id, "name": req.name}


@app.get("/api/workspaces/{ws_id}")
def api_get_workspace(ws_id: str):
    ws = get_workspace(ws_id)
    if not ws:
        raise HTTPException(404, "Workspace not found")
    return ws


@app.delete("/api/workspaces/{ws_id}")
def api_delete_workspace(ws_id: str):
    delete_workspace(ws_id)
    return {"ok": True}


# ───────────── Category API ─────────────

@app.get("/api/workspaces/{ws_id}/categories")
def api_list_categories(ws_id: str):
    return {"categories": list_categories(ws_id), "tree": get_category_tree(ws_id)}


@app.post("/api/categories")
def api_create_category(req: CategoryCreate):
    cat_id = create_category(req.name, req.workspace_id, req.parent_id, req.icon)
    if not cat_id:
        raise HTTPException(400, "Create failed")
    return {"id": cat_id}


@app.put("/api/categories/{cat_id}")
def api_update_category(cat_id: str, req: CategoryUpdate):
    if req.name or req.icon:
        if not update_category(cat_id, name=req.name, icon=req.icon):
            raise HTTPException(400, "Update failed")
    return {"ok": True}


@app.delete("/api/categories/{cat_id}")
def api_delete_category(cat_id: str):
    # 删除该分类下的所有文档
    for d in list_documents(category_id=cat_id):
        delete_document(d["doc_id"])
    for child in get_all_descendants(cat_id):
        for d in list_documents(category_id=child):
            delete_document(d["doc_id"])
    delete_category(cat_id)
    return {"ok": True}


# ───────────── Document API ─────────────

@app.get("/api/workspaces/{ws_id}/documents")
def api_list_documents(ws_id: str, category_id: Optional[str] = None, tag: Optional[str] = None, search: Optional[str] = None):
    docs = list_documents(workspace_id=ws_id, category_id=category_id, tag=tag)
    if search:
        s = search.lower()
        docs = [d for d in docs if s in d.get("filename", "").lower() or any(s in t.lower() for t in d.get("tags", []))]
    return {"documents": docs}


@app.get("/api/documents/{doc_id}")
def api_get_document(doc_id: str):
    doc = get_document(doc_id)
    if not doc:
        raise HTTPException(404, "Not found")
    # 读取内容
    try:
        content = Path(doc["filename"]).read_text(encoding="utf-8", errors="ignore")
    except Exception:
        try:
            kb = get_kb_service()
            results = kb.chroma.get(where={"doc_id": doc_id})
            content = "\n\n".join(results.get("documents", []))
        except Exception:
            content = ""
    doc["content"] = content
    return doc


@app.delete("/api/documents/{doc_id}")
def api_delete_document(doc_id: str):
    delete_document(doc_id)
    return {"ok": True}


@app.post("/api/upload")
async def api_upload(
    file: UploadFile = File(...),
    workspace_id: str = Form(...),
    category_id: Optional[str] = Form(None),
):
    file_bytes = await file.read()
    file_name = file.filename
    file_ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else "md"

    if file_ext not in ALLOWED_EXTS:
        raise HTTPException(400, f"不支持的格式: {file_ext}")

    text = extract_text(file_bytes, file_ext)
    pred_cat = auto_categorize(text)
    pred_tags = extract_tags(text, use_llm=False)

    saved = save_upload_bytes(file_bytes, file_name, file_ext)
    result = get_kb_service().upload_by_str(
        data=text, filename=saved.as_posix(),
        workspace_id=workspace_id, category_id=category_id,
        category=pred_cat, tags=pred_tags,
    )
    return {
        **result,
        "filename": file_name,
        "ext": file_ext,
        "auto_category": pred_cat,
        "auto_tags": pred_tags,
    }


@app.post("/api/notes")
def api_create_note(req: NoteCreate):
    full_content = f"# {req.title}\n\n{req.content}" if req.title else req.content
    safe = "".join(c for c in req.title if c.isalnum() or c in " _-中文")[:30] or f"note_{int(time.time())}"
    _kb = get_kb_service()
    result = _kb.upload_by_str(
        data=full_content, filename=f"{safe}.md",
        workspace_id=req.workspace_id, category_id=req.category_id,
        category="笔记", tags=req.tags or extract_tags(full_content, use_llm=False),
        source_type="note",
    )
    return result


# ───────────── Tag API ─────────────

@app.get("/api/workspaces/{ws_id}/tags")
def api_get_tags(ws_id: str):
    return {"tags": get_all_tags(workspace_id=ws_id)}


# ───────────── Chat API ─────────────

@app.post("/api/chat")
def api_chat(req: ChatRequest):
    rag = get_rag_service(req.workspace_id, req.category_id)
    try:
        answer = rag.chain.invoke(
            {"input": req.question},
            {"configurable": {"session_id": req.session_id}},
        )
        return {
            "answer": answer,
            "sources": rag.last_sources,
            "session_id": req.session_id,
        }
    except Exception as e:
        raise HTTPException(500, f"Chat failed: {e}")


@app.get("/api/chat/history")
def api_chat_history(session_id: str = "default"):
    hist = get_history(session_id)
    msgs = []
    for m in hist.messages:
        if hasattr(m, 'type') and hasattr(m, 'content'):
            msgs.append({"role": m.type, "content": m.content})
    return {"messages": msgs}


@app.delete("/api/chat/history")
def api_chat_clear(session_id: str = "default"):
    get_history(session_id).clear()
    return {"ok": True}


@app.post("/api/chat/save")
def api_save_answer(req: SaveAnswerRequest):
    result = save_answer_to_kb(
        kb_service=get_kb_service(),
        question=req.question, answer=req.answer,
        workspace_id=req.workspace_id, category_id=req.category_id,
        title=req.title, sources=req.sources,
    )
    return result


# ───────────── Stats API ─────────────

@app.get("/api/stats")
def api_stats(workspace_id: Optional[str] = None):
    if workspace_id:
        docs = list_documents(workspace_id=workspace_id)
        cats = list_categories(workspace_id)
        tags = get_all_tags(workspace_id=workspace_id)
    else:
        docs = list_documents()
        cats = []
        for ws in list_workspaces():
            cats.extend(list_categories(ws["id"]))
        tags = get_all_tags()
    return {
        "documents": len(docs),
        "categories": len(cats),
        "tags": len(tags),
        "chunks": sum(d.get("chunk_count", 0) for d in docs),
        "workspaces": len(list_workspaces()),
    }


# ───────────── Static Files (前端) ─────────────

STATIC_DIR.mkdir(exist_ok=True)

if (STATIC_DIR / "index.html").exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

    @app.get("/")
    def serve_index():
        return FileResponse(STATIC_DIR / "index.html")
else:
    @app.get("/")
    def serve_root():
        return JSONResponse({"message": "前端未构建, 请先创建 static/index.html"})


if __name__ == "__main__":
    import uvicorn
    print("=" * 50)
    print("🚀 Knowledge Base API 启动中...")
    print("📖 访问 http://localhost:8000")
    print("=" * 50)
    uvicorn.run(app, host="0.0.0.0", port=8000)
